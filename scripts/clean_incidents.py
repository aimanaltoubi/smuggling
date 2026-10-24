"""Convert local Excel/CSV incident data to validated UTF-8 CSV."""

import argparse
import csv
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
import json
from pathlib import Path
import re
import sys


REQUIRED = (
    "incident_id", "incident_date", "location_name", "governorate",
    "latitude", "longitude", "incident_type", "detained_count",
)
INCIDENT_TYPES = {"\u062a\u0647\u0631\u064a\u0628", "\u062a\u0633\u0644\u0644"}
DIGITS = str.maketrans(
    "\u0660\u0661\u0662\u0663\u0664\u0665\u0666\u0667\u0668\u0669"
    "\u06f0\u06f1\u06f2\u06f3\u06f4\u06f5\u06f6\u06f7\u06f8\u06f9",
    "01234567890123456789",
)


def text(value):
    return "" if value is None else str(value).strip()


def clean_date(value):
    if isinstance(value, (datetime, date)):
        return value.strftime("%Y-%m-%d")
    value = text(value).translate(DIGITS)
    match = re.fullmatch(r"(\d{4}[-/]\d{1,2}[-/]\d{1,2}|\d{1,2}[-/]\d{1,2}[-/]\d{4})(?:[ T].*)?", value)
    if not match:
        raise ValueError("expected a Gregorian YYYY-MM-DD or DD-MM-YYYY date (slashes also accepted)")
    parts = match[1].replace("/", "-").split("-")
    if len(parts[0]) == 4:
        year, month, day = map(int, parts)
    else:
        day, month, year = map(int, parts)
    return date(year, month, day).isoformat()


def clean_number(value, field):
    value = text(value).translate(DIGITS).replace("\u066b", ".").replace("\u066c", "")
    try:
        number = Decimal(value)
    except InvalidOperation:
        raise ValueError("expected a number; use a dot or Arabic decimal separator") from None
    if not number.is_finite():
        raise ValueError("number must be finite")
    if field == "detained_count":
        if number < 0 or number != number.to_integral_value():
            raise ValueError("must be a non-negative integer")
        if number > 9007199254740991:
            raise ValueError("exceeds JavaScript's safe integer range")
        return str(int(number))
    limit = 90 if field == "latitude" else 180
    if abs(number) > limit:
        raise ValueError(f"must be between {-limit} and {limit}")
    return format(number, "f")


def read_table(source, encoding, delimiter, sheet):
    if source.suffix.lower() == ".csv":
        if sheet:
            raise ValueError("--sheet is only supported for Excel files")
        with source.open(encoding=encoding, newline="") as handle:
            return list(csv.reader(handle, delimiter=delimiter, strict=True))
    if source.suffix.lower() != ".xlsx":
        raise ValueError("input must be .xlsx or .csv; save legacy .xls as .xlsx first")
    try:
        from openpyxl import load_workbook
    except ImportError:
        raise ValueError("Excel input requires openpyxl: py -m pip install openpyxl==3.1.5") from None
    workbook = load_workbook(source, read_only=True, data_only=False)
    try:
        worksheet = workbook[sheet] if sheet else workbook.active
        rows = []
        for cells in worksheet.iter_rows():
            for cell in cells:
                if cell.data_type in {"f", "e"}:
                    raise ValueError(f"Excel cell {cell.coordinate}: replace formulas/errors with verified values first")
            rows.append([cell.value for cell in cells])
        while rows and all(value is None for value in rows[0][-1:]):
            if not rows[0]:
                break
            if any(row and row[-1] is not None for row in rows):
                break
            rows = [row[:-1] for row in rows]
        return rows
    finally:
        workbook.close()


def clean_table(table, mapping):
    if not table:
        raise ValueError("input is empty")
    headers = [mapping.get(text(value).lstrip("\ufeff"), text(value).lstrip("\ufeff")) for value in table[0]]
    if any(not isinstance(header, str) or not header for header in headers):
        raise ValueError("headers and mapping targets must be non-empty strings")
    if len(set(headers)) != len(headers):
        raise ValueError("duplicate column names after mapping")
    missing = set(REQUIRED) - set(headers)
    if missing:
        raise ValueError("missing columns after mapping: " + ", ".join(sorted(missing)))
    cleaned, errors = [], []
    for row_number, values in enumerate(table[1:], start=2):
        if not any(text(value) for value in values):
            continue
        if len(values) != len(headers):
            errors.append(f"Row {row_number}: column count does not match the header")
            continue
        row = {header: text(value) for header, value in zip(headers, values)}
        for field, value in row.items():
            if "\ufffd" in value or "???" in value:
                errors.append(f"Row {row_number}, {field}: possible lost text; recover from the original Excel file")
        for field in REQUIRED:
            if not row[field]:
                errors.append(f"Row {row_number}, {field}: required value is missing")
        try:
            row["incident_date"] = clean_date(values[headers.index("incident_date")])
        except ValueError as error:
            errors.append(f"Row {row_number}, incident_date: {error}")
        for field in ("latitude", "longitude", "detained_count"):
            try:
                row[field] = clean_number(row[field], field)
            except ValueError as error:
                errors.append(f"Row {row_number}, {field}: {error}")
        if row["incident_type"] not in INCIDENT_TYPES:
            errors.append(f"Row {row_number}, incident_type: must be the Arabic value for smuggling or infiltration")
        cleaned.append(row)
    if not cleaned and not errors:
        errors.append("input has no incident records")
    if errors:
        raise ValueError("No output written. Fix these issues and rerun:\n" + "\n".join(errors))
    return headers, cleaned


def main(argv=None):
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("input", type=Path)
    parser.add_argument("output", type=Path)
    parser.add_argument("--column-map", type=Path, help="UTF-8 JSON object: source header to English app header")
    parser.add_argument("--sheet", help="Excel worksheet name; defaults to the active worksheet")
    parser.add_argument("--encoding", default="utf-8-sig", help="CSV input encoding, e.g. cp1256; output is always UTF-8")
    parser.add_argument("--delimiter", default=",", help="CSV input delimiter; output always uses commas")
    args = parser.parse_args(argv)
    try:
        if args.input.resolve() == args.output.resolve() or args.output.exists():
            raise ValueError("output must be a new file, different from the source")
        if args.output.suffix.lower() != ".csv":
            raise ValueError("output filename must end in .csv")
        if len(args.delimiter) != 1:
            raise ValueError("--delimiter must be one character")
        mapping = json.loads(args.column_map.read_text(encoding="utf-8-sig")) if args.column_map else {}
        if not isinstance(mapping, dict) or not all(isinstance(key, str) and isinstance(value, str) for key, value in mapping.items()):
            raise ValueError("column map must be a JSON object with string keys and values")
        table = read_table(args.input, args.encoding, args.delimiter, args.sheet)
        headers, rows = clean_table(table, mapping)
        with args.output.open("x", encoding="utf-8-sig", newline="") as handle:
            writer = csv.DictWriter(handle, fieldnames=headers)
            writer.writeheader()
            writer.writerows(rows)
        print(f"Saved {len(rows)} records as UTF-8 CSV: {args.output}")
        return 0
    except (ValueError, OSError, csv.Error, KeyError, LookupError) as error:
        print(f"Cleaning failed: {error}", file=sys.stderr)
        return 1


if __name__ == "__main__":
    sys.exit(main())