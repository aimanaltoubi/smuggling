import contextlib
import csv
from datetime import datetime
import io
import json
from pathlib import Path
import subprocess
import tempfile
import unittest

from clean_incidents import REQUIRED, clean_date, clean_table, main


class CleaningTests(unittest.TestCase):
    def record(self):
        return ["001", "23-06-2025 13:87:02", "\u0635\u062d\u0627\u0631", "\u0634\u0645\u0627\u0644 \u0627\u0644\u0628\u0627\u0637\u0646\u0629", "\u0662\u0664\u066b\u0663", "56.7", "\u062a\u0647\u0631\u064a\u0628", "\u0663"]

    def test_date_formats_and_invalid_time(self):
        for value in ("23-06-2025 13:87:02", "23/06/2025", "2025-06-23T12:00:00Z", "\u0662\u0663-\u0660\u0666-\u0662\u0660\u0662\u0665", datetime(2025, 6, 23)):
            self.assertEqual(clean_date(value), "2025-06-23")
        self.assertEqual(clean_date("03/04/2025"), "2025-04-03")
        for value in ("31-02-2025", "06-23-2025", "not a date", "45831"):
            with self.assertRaises(ValueError):
                clean_date(value)

    def test_mapping_digits_and_text_preservation(self):
        headers = list(REQUIRED) + ["details"]
        headers[1] = "\u0627\u0644\u062a\u0627\u0631\u064a\u062e"
        details = 'Arabic text, "quoted"\nnext line'
        _, rows = clean_table([headers, self.record() + [details]], {headers[1]: "incident_date"})
        self.assertEqual(rows[0]["incident_id"], "001")
        self.assertEqual(rows[0]["latitude"], "24.3")
        self.assertEqual(rows[0]["detained_count"], "3")
        self.assertEqual(rows[0]["location_name"], self.record()[2])
        self.assertEqual(rows[0]["details"], details)

    def test_invalid_rows_are_not_silently_accepted(self):
        for field, value in (("latitude", "91"), ("longitude", "NaN"), ("detained_count", "-1"), ("detained_count", "1.5"), ("incident_type", "unknown"), ("location_name", "???"), ("governorate", "")):
            record = self.record()
            record[list(REQUIRED).index(field)] = value
            with self.assertRaisesRegex(ValueError, field):
                clean_table([list(REQUIRED), record], {})
        with self.assertRaisesRegex(ValueError, "duplicate"):
            clean_table([list(REQUIRED) + ["incident_date"]], {})

    def test_csv_cli_and_app_importer(self):
        with tempfile.TemporaryDirectory() as directory:
            source, output = Path(directory) / "source.csv", Path(directory) / "clean.csv"
            with source.open("w", encoding="utf-8", newline="") as handle:
                writer = csv.writer(handle)
                writer.writerows([REQUIRED, self.record()])
            original = source.read_bytes()
            with contextlib.redirect_stdout(io.StringIO()):
                self.assertEqual(main([str(source), str(output)]), 0)
            self.assertEqual(source.read_bytes(), original)
            self.assertTrue(output.read_bytes().startswith(b"\xef\xbb\xbf"))
            self.assertIn(self.record()[2], output.read_text(encoding="utf-8-sig"))
            result = subprocess.run([
                "node", "--input-type=module", "-e",
                "import { readFileSync } from 'node:fs'; "
                "import { parseIncidentCsv } from './src/lib/analytics/csvImport.js'; "
                "const result = await parseIncidentCsv({text: async () => readFileSync(process.argv[1], 'utf8')}); "
                "console.log(JSON.stringify(result));", str(output),
            ], cwd=Path(__file__).resolve().parents[1], capture_output=True, text=True, check=True)
            result = json.loads(result.stdout)
            self.assertEqual(result["errors"], [])
            self.assertEqual(len(result["rows"]), 1)
            with contextlib.redirect_stderr(io.StringIO()):
                self.assertEqual(main([str(source), str(output)]), 1)

    def test_invalid_input_does_not_write_output(self):
        with tempfile.TemporaryDirectory() as directory:
            source, output = Path(directory) / "source.csv", Path(directory) / "clean.csv"
            source.write_text(",".join(REQUIRED) + "\n" + ",".join(self.record()).replace("23-06-2025", "31-02-2025"), encoding="utf-8")
            with contextlib.redirect_stderr(io.StringIO()):
                self.assertEqual(main([str(source), str(output)]), 1)
            self.assertFalse(output.exists())

    def test_legacy_csv_encoding_delimiter_and_mapping(self):
        with tempfile.TemporaryDirectory() as directory:
            source, output = Path(directory) / "source.csv", Path(directory) / "clean.csv"
            mapping = Path(directory) / "columns.json"
            headers = list(REQUIRED)
            headers[1] = "\u0627\u0644\u062a\u0627\u0631\u064a\u062e"
            mapping.write_text(json.dumps({headers[1]: "incident_date"}), encoding="utf-8")
            record = self.record()
            record[4], record[7] = "24.3", "3"
            with source.open("w", encoding="cp1256", newline="") as handle:
                csv.writer(handle, delimiter=";").writerows([headers, record])
            with contextlib.redirect_stdout(io.StringIO()):
                self.assertEqual(main([str(source), str(output), "--encoding", "cp1256", "--delimiter", ";", "--column-map", str(mapping)]), 0)
            with output.open(encoding="utf-8-sig", newline="") as handle:
                rows = list(csv.DictReader(handle))
            self.assertEqual(rows[0]["location_name"], record[2])
            self.assertEqual(rows[0]["incident_date"], "2025-06-23")

    def test_excel_conversion(self):
        try:
            from openpyxl import Workbook
        except ImportError:
            self.skipTest("openpyxl is not installed")
        with tempfile.TemporaryDirectory() as directory:
            source, output = Path(directory) / "source.xlsx", Path(directory) / "clean.csv"
            workbook = Workbook()
            workbook.active.append(list(REQUIRED))
            record = self.record()
            record[1] = datetime(2025, 6, 23, 13, 2)
            workbook.active.append(record)
            workbook.save(source)
            with contextlib.redirect_stdout(io.StringIO()):
                self.assertEqual(main([str(source), str(output)]), 0)
            with output.open(encoding="utf-8-sig", newline="") as handle:
                rows = list(csv.DictReader(handle))
            self.assertEqual(rows[0]["incident_date"], "2025-06-23")
            self.assertEqual(rows[0]["location_name"], record[2])

    def test_excel_sheet_selection_and_formula_rejection(self):
        try:
            from openpyxl import Workbook
        except ImportError:
            self.skipTest("openpyxl is not installed")
        with tempfile.TemporaryDirectory() as directory:
            source, output = Path(directory) / "source.xlsx", Path(directory) / "clean.csv"
            workbook = Workbook()
            worksheet = workbook.create_sheet("Incidents")
            worksheet.append(list(REQUIRED))
            worksheet.append(self.record())
            workbook.save(source)
            with contextlib.redirect_stdout(io.StringIO()):
                self.assertEqual(main([str(source), str(output), "--sheet", "Incidents"]), 0)
            worksheet["H2"] = "=1+2"
            workbook.save(source)
            output = output.with_name("formula.csv")
            with contextlib.redirect_stderr(io.StringIO()) as errors:
                self.assertEqual(main([str(source), str(output), "--sheet", "Incidents"]), 1)
            self.assertIn("H2", errors.getvalue())
            self.assertFalse(output.exists())


if __name__ == "__main__":
    unittest.main()